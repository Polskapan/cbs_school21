#!/usr/bin/env python3
"""
Проверка корректности инвентаризации ИТ-активов (My_inventory.xlsx).

Правила:
  1. У каждого устройства указан хотя бы один IPv4-адрес хоста
     (маски подсети вида 255.255.255.0 за адрес не считаются).
  2. Указана версия ОС (столбец «ОС с версией» не пуст и содержит номер версии).
  3. У НЕ-сетевых устройств указано хотя бы одно ПО.
     Сетевое оборудование (роутеры, коммутаторы, файрволы, точки доступа)
     от требования ПО освобождается.

Использование:
    python validate_inventory.py [путь_к_файлу.xlsx]

Код возврата: 0 — все проверки пройдены, 1 — есть нарушения,
2 — файл/шапку прочитать не удалось (удобно для pre-commit / CI).
"""

import re
import sys
import ipaddress
from pathlib import Path

import openpyxl

DEFAULT_FILE = "My_inventory.xlsx"

# Заголовки столбцов ищем по названию, а не по номеру — устойчиво к сдвигам.
H_DEVICE   = "Устройство"
H_OS       = "ОС с версией"
H_SOFTWARE = "Установленное ПО"   # по префиксу: в шапке "... (минимум 1)"
H_IP       = "IP-адрес"           # по префиксу: в шапке "IP-адрес/маска"

# Признаки сетевого оборудования — для таких устройств ПО не обязательно.
# При необходимости дополните список (например, добавьте "принтер"/"ибп",
# если по вашей методичке они тоже освобождаются от требования ПО).
NETWORK_KEYWORDS = (
    "cisco", "catalyst", "nexus", "isr", "asr",
    "коммутатор", "маршрутизатор", "роутер",
    "switch", "router", "firewall",
    "межсетев", "файрвол", "фаервол", "брандмауэр",
    "точка доступа", "access point", "wlc", "контроллер"
)

IPV4_RE = re.compile(r"\b\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3}\b")


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def is_contiguous_netmask(addr: ipaddress.IPv4Address) -> bool:
    """True, если адрес — это маска подсети (255.255.255.0 и т.п.), а не IP хоста."""
    val = int(addr)
    if val == 0:
        return False
    inverted = val ^ 0xFFFFFFFF          # маска = единицы, затем нули
    return (inverted & (inverted + 1)) == 0


def host_ipv4_addresses(text: str) -> list:
    """Валидные IPv4-адреса хоста из текста; маски подсети отбрасываются."""
    found = []
    for candidate in IPV4_RE.findall(text or ""):
        try:
            addr = ipaddress.IPv4Address(candidate)
        except ipaddress.AddressValueError:
            continue                     # например 999.1.1.1 — не валидный IPv4
        if is_contiguous_netmask(addr):
            continue
        found.append(str(addr))
    return found


def has_os_version(text: str) -> bool:
    """ОС не пуста и содержит номер версии (хотя бы одну цифру)."""
    return (not is_blank(text)) and any(ch.isdigit() for ch in str(text))


def is_network_device(name: str) -> bool:
    low = (name or "").lower()
    return any(kw in low for kw in NETWORK_KEYWORDS)


def find_header(ws):
    """Находит строку заголовков и сопоставляет нужные столбцы по названию."""
    for row in ws.iter_rows(min_row=1, max_row=30):
        texts = {c.column: (str(c.value).strip() if c.value is not None else "")
                 for c in row}
        if H_DEVICE not in texts.values():
            continue
        cols = {}
        for col, text in texts.items():
            if text == H_DEVICE:
                cols["device"] = col
            elif text == H_OS:
                cols["os"] = col
            elif text.startswith(H_SOFTWARE):
                cols["software"] = col
            elif text.startswith(H_IP):
                cols["ip"] = col
        return row[0].row, cols
    return None, None


def validate(path: str) -> int:
    p = Path(path)
    if not p.exists():
        print(f"[ОШИБКА] Файл не найден: {p}")
        return 2

    wb = openpyxl.load_workbook(p)
    ws = wb.active

    header_row, cols = find_header(ws)
    if not header_row:
        print(f"[ОШИБКА] Не найдена строка заголовков со столбцом «{H_DEVICE}».")
        return 2

    missing = [n for n in ("device", "os", "software", "ip") if n not in cols]
    if missing:
        print(f"[ОШИБКА] В шапке не найдены столбцы: {missing}")
        return 2

    total = failed = 0

    for row in ws.iter_rows(min_row=header_row + 1):
        row_no = row[0].row
        get = lambda key: ws.cell(row=row_no, column=cols[key]).value

        if is_blank(get("device")):
            continue                     # пустая строка-заготовка — пропускаем

        total += 1
        device = str(get("device")).splitlines()[0].strip()
        problems = []

        # 1. Хотя бы один IPv4-адрес хоста
        if not host_ipv4_addresses(str(get("ip") or "")):
            problems.append("нет валидного IPv4-адреса")

        # 2. Версия ОС
        if not has_os_version(str(get("os") or "")):
            problems.append("не указана версия ОС")

        # 3. ПО для не-сетевых устройств
        net = is_network_device(str(get("device")))
        if not net and is_blank(get("software")):
            problems.append("нет ни одного ПО")

        if problems:
            failed += 1
            print(f"[FAIL] стр.{row_no:>3} | {device}")
            for pr in problems:
                print(f"         └─ {pr}")
        else:
            tag = "  (сетевое устройство — ПО не требуется)" if net else ""
            print(f"[ OK ] стр.{row_no:>3} | {device}{tag}")

    print("-" * 62)
    print(f"Проверено устройств: {total} | Ошибок: {failed}")
    return 1 if failed else 0


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FILE
    sys.exit(validate(target))
